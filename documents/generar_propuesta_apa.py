#!/usr/bin/env python3
"""
Genera la propuesta de ScoreUp en formato APA (7a ed.) como archivo .docx
"""
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import os

doc = Document()

# ── Estilos globales APA ──────────────────────────────────────────────
style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)
pf = style.paragraph_format
pf.line_spacing = 2.0
pf.space_after = Pt(0)
pf.space_before = Pt(0)

for section in doc.sections:
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)

# ── Helpers ───────────────────────────────────────────────────────────
def add_title_page():
    for _ in range(6):
        doc.add_paragraph()
    add_centered('TALLER DE PROGRAMACIÓN', bold=True)
    doc.add_paragraph()
    add_centered('PROPUESTA INICIAL', bold=True)
    doc.add_paragraph()
    add_centered('ScoreUp: Sistema de Evaluación de Clubes de Conquistadores', bold=True)
    for _ in range(3):
        doc.add_paragraph()
    add_centered('Integrantes:')
    doc.add_paragraph()
    for name in [
        'JUAN FERNANDO NINA CACHI',
        'AXEL ALBERTO TICONA VILLEGAS',
        'KEVIN QUISPE CANAVIRI',
        'ERICK LANCHIMBA LANCHIMBA',
    ]:
        add_centered(name)
    for _ in range(3):
        doc.add_paragraph()
    add_centered('Carrera: Ingeniería de Sistemas — 4to Semestre')
    add_centered('Materia: Taller de Programación')
    doc.add_page_break()

def add_centered(text, bold=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = bold
    return p

def h1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(12)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(14)
    return p

def h2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(12)
    return p

def h3(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(12)
    return p

def para(text, bold=False, italic=False, indent=False):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Cm(1.27)
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    return p

def bullet(text, bold_prefix=''):
    p = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text)
    else:
        p.add_run(text)
    return p

def add_table(headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = ''
        p = hdr_cells[i].paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.size = Pt(10)
        run.font.name = 'Times New Roman'
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="D9E2F3"/>')
        hdr_cells[i]._tc.get_or_add_tcPr().append(shading)
    for row_data in rows:
        row_cells = table.add_row().cells
        for i, cell_text in enumerate(row_data):
            row_cells[i].text = ''
            p = row_cells[i].paragraphs[0]
            run = p.add_run(cell_text)
            run.font.size = Pt(10)
            run.font.name = 'Times New Roman'
    return table

# ── Portada ───────────────────────────────────────────────────────────
add_title_page()

# ── 1. Descripción del problema ───────────────────────────────────────
h1('1. Descripción del problema')
para(
    'En una organización de Conquistadores se realizan periódicamente eventos '
    '(campamentos, reuniones, actividades regionales) en los que participan varios clubes. '
    'Actualmente, la evaluación del desempeño de cada club se maneja de forma manual: planillas '
    'en papel u hojas de cálculo, criterios definidos de manera informal y resultados que no '
    'siempre pueden rastrearse a una evaluación original.', indent=True
)
para('Esta situación genera varios problemas concretos:', indent=True)
bullet('La información está dispersa entre planillas, correos y mensajes.')
bullet(
    'No se puede responder con certeza preguntas como: ¿qué criterios se evaluaron?, '
    '¿quién evaluó a cada club?, ¿cuál fue el porcentaje ponderado?, ¿qué club ganó la temporada?, '
    '¿qué penalizaciones se aplicaron y por qué motivo?'
)
bullet('Los cálculos (porcentaje ponderado, desempates, suma de penalizaciones) se realizan a mano y son propensos a errores.')
bullet(
    'No existe una forma confiable de cerrar un evento: se puede evaluar aún cuando los criterios '
    'no están completos o sus pesos no suman 100%.'
)
bullet('Los incidentes disciplinarios quedan solo como relato verbal, sin registro con evidencia.')

# ── 2. Problema central ──────────────────────────────────────────────
h1('2. Problema central a investigar')
para(
    '¿Cómo mejorar el registro, cálculo y seguimiento de las evaluaciones de los clubes de '
    'Conquistadores durante sus eventos y temporadas mediante una aplicación informática que '
    'centralice la información, automatice los cálculos y permita la consulta de resultados por '
    'parte de todos los actores (administradores, evaluadores y directores de club)?', italic=True, indent=True
)
para(
    'El problema es relevante porque el resultado de la evaluación determina el ranking de cada '
    'temporada y tiene un impacto directo en la motivación y reconocimiento de los clubes. Un error '
    'de cálculo o un dato perdido puede cambiar un resultado oficial.', indent=True
)

# ── 3. Justificación ─────────────────────────────────────────────────
h1('3. Justificación')
bullet('Automatización de cálculos: ', bold_prefix='Automatización de cálculos: ')
bullet(
    'el porcentaje ponderado, el puntaje por evento y el ranking de temporada se calculan de forma '
    'sistemática, eliminando errores manuales.'
)
bullet(
    'Centralización: toda la información de clubes, eventos, evaluaciones, penalizaciones e '
    'incidentes queda en un único sistema con persistencia en SQLite.'
)
bullet(
    'Transparencia: cada evaluación queda asociada a un evaluador (staff principal) y a una fecha, '
    'lo que permite auditar los resultados.'
)
bullet(
    'Control: se impone una regla de negocio que impide cerrar un evento si los pesos de los '
    'criterios no suman 100%.'
)
bullet(
    'Escalabilidad: al separar dominio, cálculo y persistencia, el sistema puede crecer '
    '(nuevos tipos de evaluación, nuevos reportes) sin reescribir componentes enteros.'
)

# ── 4. Objetivo general ──────────────────────────────────────────────
h1('4. Objetivo general')
para(
    'Desarrollar un sistema informático (ScoreUp) que permita a una organización de Conquistadores '
    'registrar, evaluar, calificar y dar seguimiento al desempeño de sus clubes en eventos y '
    'temporadas, con resultados confiables y consultables por todos los actores involucrados.', indent=True
)

# ── 5. Objetivos específicos ─────────────────────────────────────────
h1('5. Objetivos específicos')
objs = [
    ('OE-01. ', 'Identificar los criterios e información necesarios para evaluar el desempeño de un club en un evento (criterios con peso, puntaje máximo).'),
    ('OE-02. ', 'Diseñar un mecanismo para registrar evaluaciones por club y por evento, con valores por criterio y comentario del evaluador.'),
    ('OE-03. ', 'Implementar los cálculos de porcentaje ponderado, puntaje por evento, desempates y ranking de temporada, considerando penalizaciones.'),
    ('OE-04. ', 'Permitir la consulta de resultados y reportes por parte de administradores y directores de club, con exportación a Excel y PDF.'),
    ('OE-05. ', 'Diseñar una arquitectura de software modular (interfaz, lógica de aplicación, dominio, acceso a datos) que facilite el mantenimiento y la evolución del sistema.'),
]
for i, (code, desc) in enumerate(objs, 1):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1.27)
    r = p.add_run(f'{code}')
    r.bold = True
    p.add_run(desc)

# ── 6. Alcance ───────────────────────────────────────────────────────
h1('6. Alcance')
h2('6.1 Funcionalidades mínimas')
funcs = [
    ('A. Gestión de la organización y usuarios',
     ['Registro de una organización con su usuario administrador inicial.',
      'Roles con permisos: Administrador, Staff (evaluador) y Director de club.',
      'Creación, edición y eliminación de usuarios; asignación de contraseñas seguras.']),
    ('B. Gestión de clubes y miembros',
     ['Alta, baja y modificación de clubes.',
      'Registro de miembros (nombre, edad, cargo, año de ingreso, categoría).',
      'Importación de miembros desde archivo Excel.']),
    ('C. Gestión de eventos',
     ['Creación de eventos (fecha, lugar, tipo, temporada, puntaje máximo).',
      'Configuración de criterios con peso (la suma debe alcanzar 100% para cerrar).',
      'Asignación de clubes participantes y de un staff principal por club.',
      'Cierre del evento (bloquea modificaciones y evaluaciones).']),
    ('D. Evaluación',
     ['Registro de evaluaciones por staff: valor 0–100 por criterio y comentario.',
      'Identificación de la evaluación oficial (la del staff principal) y respaldos.',
      'Cálculo automático de porcentaje ponderado y puntaje.']),
    ('E. Penalizaciones e incidentes',
     ['Catálogo de motivos de penalización (descripción y puntos de descuento).',
      'Aplicación y anulación de penalizaciones por staff.',
      'Registro de incidentes con tipo, gravedad, descripción y evidencia fotográfica.']),
    ('F. Resultados y reportes',
     ['Ranking por evento y por temporada (puntaje bruto − penalizaciones).',
      'Regla de desempate por el criterio de mayor peso.',
      'Reporte de actividad del staff (clubes asignados vs. evaluaciones completadas).',
      'Exportación de resultados a Excel y PDF.']),
]
for title, items in funcs:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1.27)
    r = p.add_run(title)
    r.bold = True
    for item in items:
        bullet(item)

h2('6.2 Fuera del alcance de la primera versión')
oos = [
    'Aplicación móvil nativa.',
    'Notificaciones automáticas por correo o mensajería.',
    'Integración con sistemas institucionales externos.',
    'Módulo de pagos o inscripciones.',
    'Sistema de autoevaluación por parte de los clubes.',
    'Predicción de desempeño mediante Machine Learning.',
]
for item in oos:
    bullet(item)
para('Estas funcionalidades pueden plantearse como trabajo futuro.', indent=True)

# ── 7. Identificación de actores ─────────────────────────────────────
h1('7. Identificación de actores')
add_table(
    ['Actor', 'Responsabilidad'],
    [
        ['Administrador',
         'Gestiona la organización, los clubes, los eventos, los criterios y los usuarios. '
         'Genera accesos masivos y exporta resultados.'],
        ['Staff evaluador',
         'Evalúa a los clubes asignados, registra incidentes y aplica penalizaciones. '
         'Es el evaluador principal u oficial de un club en el evento.'],
        ['Director de club',
         'Consulta el ranking de temporada y los puntajes obtenidos por su club en los eventos cerrados.'],
    ]
)
para('')
para(
    'Se mantienen los tres roles porque son los que aparecen en el flujo real del problema: '
    'quien configura y administra (admin), quien evalúa (staff) y quien consume resultados (director). '
    'No se incorpora un rol de "ciudadano" porque el contexto no lo requiere; podrá agregarse en el '
    'futuro si la organización lo solicita.', indent=True
)

# ── 8. Requerimientos funcionales ────────────────────────────────────
h1('8. Requerimientos funcionales')
rfs = [
    ('RF-01', 'El sistema deberá permitir registrar una nueva organización junto con su usuario administrador inicial.'),
    ('RF-02', 'El sistema deberá autenticar a los usuarios mediante usuario y contraseña y validar sus permisos según el rol (admin, staff, director).'),
    ('RF-03', 'El sistema deberá permitir crear, renombrar y eliminar clubes.'),
    ('RF-04', 'El sistema deberá permitir registrar y eliminar miembros de un club, así como importarlos desde un archivo Excel.'),
    ('RF-05', 'El sistema deberá permitir crear eventos con fecha, lugar, tipo, temporada y puntaje máximo.'),
    ('RF-06', 'El sistema deberá permitir configurar criterios de evaluación con su peso por evento.'),
    ('RF-07', 'El sistema deberá impedir cerrar un evento cuando la suma de pesos de sus criterios no sea 100%.'),
    ('RF-08', 'El sistema deberá permitir asignar clubes a un evento e indicar el staff principal evaluador de cada club.'),
    ('RF-09', 'El sistema deberá permitir al staff registrar evaluaciones (valor 0–100 por criterio y comentario) para los clubes asignados.'),
    ('RF-10', 'El sistema deberá calcular automáticamente el porcentaje ponderado y el puntaje de cada evaluación por evento.'),
    ('RF-11', 'El sistema deberá calcular el ranking por evento con desempate según el criterio de mayor peso.'),
    ('RF-12', 'El sistema deberá calcular el ranking de temporada restando las penalizaciones del puntaje bruto.'),
    ('RF-13', 'El sistema deberá permitir gestionar motivos de penalización (nombre, descripción, puntos, activo/inactivo).'),
    ('RF-14', 'El sistema deberá permitir al staff aplicar y al administrador anular penalizaciones a los clubes.'),
    ('RF-15', 'El sistema deberá permitir registrar incidentes con tipo, gravedad, descripción y evidencia fotográfica.'),
    ('RF-16', 'El sistema deberá permitir filtrar incidentes por evento, club, tipo, gravedad y estado.'),
    ('RF-17', 'El sistema deberá generar masivamente usuarios y contraseñas para staff y directores de un evento.'),
    ('RF-18', 'El sistema deberá permitir importar usuarios (staff o directores) desde archivos Excel.'),
    ('RF-19', 'El sistema deberá permitir exportar los resultados a Excel y PDF.'),
    ('RF-20', 'El sistema deberá mostrar al director el puntaje y la posición de su club en la temporada y en los eventos cerrados.'),
]
add_table(['Código', 'Requerimiento'], rfs)

# ── 9. Requerimientos no funcionales ─────────────────────────────────
h1('9. Requerimientos no funcionales')
rnfs = [
    ('RNF-01', 'Usabilidad', 'La interfaz deberá ser sencilla y guiar al usuario con formularios y tablas claras; no se requiere conocimientos técnicos previos.'),
    ('RNF-02', 'Seguridad', 'El sistema deberá controlar el acceso por rol y almacenar las contraseñas de forma segura (hash con werkzeug).'),
    ('RNF-03', 'Rendimiento', 'Las consultas habituales (listados, rankings, reportes) deberán ejecutarse en un tiempo razonable para el volumen de datos esperado.'),
    ('RNF-04', 'Mantenibilidad', 'La solución deberá organizarse en capas (interfaz, lógica, dominio, persistencia) que faciliten su mantenimiento.'),
    ('RNF-05', 'Escalabilidad', 'La arquitectura deberá permitir incorporar nuevas funcionalidades (reportes, módulos) sin modificar grandes partes del sistema.'),
    ('RNF-06', 'Portabilidad', 'El sistema deberá ejecutarse con dependencias mínimas (Python + Flask + SQLite + Tailwind CSS) y ser replicable mediante un listado de requerimientos.'),
]
add_table(['Código', 'Categoría', 'Requerimiento'], rnfs)

# ── 10. Historias de usuario ─────────────────────────────────────────
h1('10. Historias de usuario')

hus = [
    ('HU-01 – Registrar un evento con criterios',
     'administrador', 'crear un evento y definir sus criterios con su peso',
     'preparar la evaluación de los clubes participantes.',
     ['Se debe indicar nombre, fecha, lugar, tipo, temporada y puntaje máximo.',
      'Se deben agregar criterios con peso en porcentaje.',
      'El sistema debe advertir cuando la suma de pesos supere 100%.',
      'El evento no se puede cerrar hasta que la suma de pesos sea 100%.']),
    ('HU-02 – Evaluar a un club',
     'staff evaluador', 'registrar la evaluación de un club asignado indicando un valor por cada criterio',
     'que su puntaje se calcule de forma automática y trazable.',
     ['Cada valor debe estar entre 0 y 100.',
      'Se debe poder incluir un comentario.',
      'Si ya existe una evaluación mía para el club, debe actualizarse (no duplicarse).',
      'El sistema debe mostrar el porcentaje ponderado y el puntaje resultante.']),
    ('HU-03 – Consultar el ranking de temporada',
     'administrador', 'ver el ranking de la temporada indicada',
     'conocer el orden final de los clubes y poder publicarlo.',
     ['El ranking debe ordenar por puntaje bruto − penalizaciones.',
      'Las posiciones deben respetar los empates (posición compartida).',
      'Debe incluir el detalle por evento y el monto de penalizaciones.']),
    ('HU-04 – Ver el desempeño de mi club',
     'director de club', 'consultar el puntaje y la posición de mi club',
     'informar a mis miembros sobre los resultados.',
     ['Solo debo ver los datos de mi propio club.',
      'Debe mostrar la posición en la temporada y los puntajes de los eventos cerrados.']),
    ('HU-05 – Reportar un incidente con evidencia',
     'staff evaluador', 'reportar un incidente disciplinario adjuntando una fotografía',
     'dejar un registro verificable del hecho.',
     ['Debo seleccionar tipo, gravedad y descripción.',
      'Debo poder adjuntar una imagen en formato permitido (png, jpg, gif, webp).',
      'El sistema debe asignar fecha y hora automáticamente.']),
    ('HU-06 – Exportar resultados',
     'administrador', 'exportar los resultados del ranking a Excel o PDF',
     'generar los documentos oficiales de la organización.',
     ['Se debe elegir formato (Excel o PDF) y alcance (temporada o evento).',
      'El archivo debe incluir encabezados, valores por criterio y totales.',
      'El archivo se debe descargar con un nombre identificable.']),
]

for title, actor, want, because, criteria in hus:
    h3(title)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1.27)
    r = p.add_run('Como ')
    r = p.add_run(actor)
    r.bold = True
    p.add_run(', quiero ')
    r2 = p.add_run(want)
    r2.bold = True
    p.add_run(', para ')
    r3 = p.add_run(because)
    r3.bold = True
    p.add_run('.')
    p2 = doc.add_paragraph()
    p2.paragraph_format.left_indent = Cm(1.27)
    r = p2.add_run('Criterios de aceptación:')
    r.bold = True
    for c in criteria:
        bullet(c)

# ── 11. Arquitectura ─────────────────────────────────────────────────
h1('11. Propuesta inicial de arquitectura')
para(
    'La solución es una arquitectura monolítica en capas, compatible con los conceptos de Clean '
    'Architecture: las dependencias apuntan hacia adentro (dominio) y la infraestructura es '
    'reemplazable.', indent=True
)
para('')
# Diagrama de arquitectura (texto)
arch_lines = [
    '    Administrador      Staff evaluador      Director de club',
    '          │                  │                     │',
    '          └──────────────┬───┴─────────────────────┘',
    '                         ▼',
    '              ┌───────────────────┐',
    '              │   Interfaz web    │   templates/ (Jinja2) + Tailwind CSS',
    '              │  (presentación)   │   + app.py (plantillas y rutas)',
    '              └──────────┬────────┘',
    '                         ▼',
    '              ┌───────────────────┐',
    '              │ Lógica de         │   app.py (rutas, control de',
    '              │  aplicación /     │   sesión, casos de uso)',
    '              │  casos de uso     │',
    '              └──────────┬────────┘',
    '                         ▼',
    '              ┌───────────────────┐',
    '              │      Dominio      │   data.py (entidades) +',
    '              │  reglas y cálculos│   calc.py (lógica de negocio)',
    '              └──────────┬────────┘',
    '                         ▼',
    '              ┌───────────────────┐',
    '              │ Acceso a datos    │   persist.py (repositorio)',
    '              │  (repositorio)    │',
    '              └──────────┬────────┘',
    '                         ▼',
    '              ┌───────────────────┐',
    '              │   SQLite  .db     │',
    '              └───────────────────┘',
]
for line in arch_lines:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(line)
    r.font.name = 'Courier New'
    r.font.size = Pt(9)

para('')
h2('Justificación de la decisión')
bullet(
    'Monolítica: el dominio es pequeño y está bien delimitado (evaluación de clubes). '
    'Una solución monolítica en capas es más simple de desarrollar, probar y mantener.'
)
bullet(
    'Separación de responsabilidades: la presentación (templates/, rutas en app.py), las entidades '
    '(data.py), la lógica de cálculo (calc.py) y la persistencia (persist.py) están separadas.'
)
bullet(
    'Facilidad de migración futura: si la organización crece o se requiere una API o aplicación '
    'móvil, la lógica de negocio de calc.py puede reutilizarse sin cambios.'
)

# ── 12. SOLID ─────────────────────────────────────────────────────────
h1('12. Aplicación de principios SOLID')

h2('12.1 Single Responsibility Principle (SRP)')
para(
    'Cada módulo tiene una única responsabilidad: data.py solo define las entidades del dominio '
    '(dataclasses); calc.py solo contiene la lógica de cálculo (ponderación, rankings, desempates); '
    'persist.py solo serializa y persiste el estado; app.py solo atiende peticiones HTTP y coordina '
    'las demás capas.', indent=True
)
para(
    'Una entidad como Evaluacion no se encarga de calcularse ni guardarse a sí misma; esas '
    'responsabilidades viven en calc.py y persist.py.', indent=True
)

h2('12.2 Open/Closed Principle (OCP)')
para(
    'El sistema está diseñado para extenderse sin modificar el núcleo: los criterios de un evento '
    'son datos configurables (se agregan nuevos criterios sin tocar el modelo); los motivos de '
    'penalización se gestionan como catálogo (un nuevo motivo no exige cambios en calc.py); el '
    'cálculo de rankings itera sobre criterios y evaluaciones genéricas.', indent=True
)

h2('12.3 Dependency Inversion Principle (DIP)')
para(
    'La lógica de dominio no depende de la base de datos concreta: calc.py trabaja sobre los '
    'objetos del dominio (listas de evaluaciones, criterios, eventos) y no conoce SQLite; persist.py '
    'actúa como intermediario.', indent=True
)
para('')
# Diagrama DIP
dip_lines = [
    '   Caso de uso (app.py)',
    '        │',
    '        ▼',
    '   persist.save / persist.load     ← interfaz de persistencia',
    '        ▲',
    '        │',
    '   Implementación SQLite',
]
for line in dip_lines:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(line)
    r.font.name = 'Courier New'
    r.font.size = Pt(10)

# ── 13. Patrones de diseño ───────────────────────────────────────────
h1('13. Patrones de diseño propuestos')

h2('13.1 Patrón Repository (Acceso a datos)')
para(
    'persist.py implementa un repositorio: el resto del sistema manipula objetos de dominio '
    '(clubs, eventos, evaluaciones) y delega en el repositorio la tarea de guardarlos y '
    'recuperarlos (es un Adapter sobre SQLite). Esto aísla el resto de la aplicación de los '
    'detalles de la base de datos y justifica el principio DIP.', indent=True
)
para(
    'Ventaja: cambios de motor de base de datos, esquema o serialización solo afectan a persist.py.', indent=True
)

h2('13.2 Patrón Factory (Creación de usuarios y entidades)')
para(
    'La creación de usuarios según rol (admin, staff, director) aparece repetida en app.py. '
    'Se propone centralizar esta creación en una fábrica de usuarios:', indent=True
)
factory_lines = [
    '             UsuarioFactory',
    '                  │',
    '      ┌───────────┼──────────────┐',
    '      ▼           ▼              ▼',
    '  Usuario      Usuario        Usuario',
    '   Admin        Staff          Director',
    ' (org)        (evento)         (club)',
]
for line in factory_lines:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(line)
    r.font.name = 'Courier New'
    r.font.size = Pt(10)
para('')
para(
    'Ventaja: se evita la repetición del código de construcción, se garantiza que cada tipo de '
    'usuario se inicialice con los campos correctos y se simplifica la generación masiva de accesos.', indent=True
)

# ── 14. DDD ───────────────────────────────────────────────────────────
h1('14. Consideración de DDD')
bullet('Dominio: evaluación y calificación del desempeño de clubes de Conquistadores en eventos y temporadas.')
bullet('Entidad principal: Club (identidad persistente: nombre, miembros, puntajes por temporada).')
bullet(
    'Entidades: Organizacion, Usuario, Miembro, Evento, Criterio, Evaluacion, '
    'Incidente, PenalizacionAplicada, MotivoPenalizacion.'
)
bullet(
    'Objetos de valor: PorcentajePonderado, Puntaje, Temporada, EstadoEvento, '
    'EstadoIncidente, Gravedad, Fecha.'
)
h3('Estados del evento')
state_ev = ['ABIERTO → (suma de pesos = 100% y cierre) → CERRADO']
for s in state_ev:
    bullet(s)
h3('Estados del incidente')
state_inc = ['ABIERTO → EN_REVISION → RESUELTO']
for s in state_inc:
    bullet(s)
h3('Reglas de negocio identificadas')
reglas = [
    'La suma de los pesos de los criterios de un evento debe ser 100% para cerrarlo.',
    'La evaluación oficial de un club es la del staff principal asignado.',
    'El desempate entre clubes se resuelve comparando el criterio de mayor peso.',
    'El puntaje de temporada es la suma de puntajes por evento menos las penalizaciones.',
]
for r in reglas:
    bullet(r)

# ── 15. Prototipo ─────────────────────────────────────────────────────
h1('15. Prototipo')
para(
    'Se cuenta con un prototipo funcional implementado con Flask, plantillas Jinja2 y Tailwind CSS '
    '(vía CDN) en templates/. Las pantallas principales son:', indent=True
)
pantallas = [
    ('Inicio / Login', 'index.html, login.html', 'Autenticación por usuario y contraseña, redirección según rol'),
    ('Registro', 'registro.html', 'Alta de organización + administrador inicial'),
    ('Dashboard admin', 'admin/dashboard.html', 'Indicadores: eventos activos/próximos, miembros, staff'),
    ('Clubes', 'admin/clubes.html, club_miembros.html', 'CRUD de clubes y miembros, importación Excel'),
    ('Eventos', 'admin/eventos.html, evento_detalle.html', 'Creación, criterios, asignación, generación de accesos'),
    ('Incidentes', 'admin/incidentes.html', 'Registro, filtros, cambio de estado y carga de evidencia fotográfica'),
    ('Penalizaciones', 'admin/motivos.html, penalizaciones.html', 'Catálogo de motivos y aplicación de descuentos'),
    ('Resultados', 'admin/resultados.html, reportes.html', 'Rankings y exportación Excel/PDF'),
    ('Panel staff', 'staff/panel.html', 'Evaluación por criterio, incidentes y penalizaciones'),
    ('Panel director', 'director/panel.html', 'Posición y puntajes de su club'),
]
add_table(['Pantalla', 'Ruta / Archivo', 'Funcionalidad'], pantallas)
para('')
para(
    'El prototipo valida el flujo de interacción end-to-end: desde la creación del evento hasta '
    'la exportación del ranking, pasando por la evaluación y las penalizaciones.', indent=True
)

# ── 16. Planificación ────────────────────────────────────────────────
h1('16. Planificación inicial')
etapas = [
    ('1', 'Análisis del problema', 'Problema y contexto'),
    ('2', 'Levantamiento de requerimientos', 'Requerimientos e historias de usuario'),
    ('3', 'Diseño arquitectónico', 'Arquitectura propuesta'),
    ('4', 'Diseño del dominio', 'Entidades y reglas de negocio'),
    ('5', 'Diseño de interfaz', 'Prototipo'),
    ('6', 'Implementación inicial', 'Primer incremento funcional'),
    ('7', 'Integración', 'Sistema integrado'),
    ('8', 'Pruebas', 'Evidencias de pruebas'),
    ('9', 'Refactorización', 'Versión mejorada'),
    ('10', 'Presentación final', 'Producto terminado'),
]
add_table(['Etapa', 'Actividad', 'Producto'], etapas)
para('')
para(
    'El proyecto ya cuenta con un primer incremento implementado (módulos de clubes, eventos, '
    'criterios, evaluaciones, rankings y exportación), de modo que las etapas 1 a 8 corresponden a '
    'consolidarlas, documentarlas y refinar la solución.', indent=True
)

# ── 17. Conclusiones ─────────────────────────────────────────────────
h1('17. Conclusiones')
conclusiones = [
    'El problema de la evaluación manual de clubes es real y genera errores y falta de trazabilidad; ScoreUp lo aborda centralizando la información y automatizando los cálculos.',
    'La arquitectura en capas (interfaz → casos de uso → dominio → repositorio) permite un mantenimiento simple y deja abierta la evolución futura del sistema.',
    'Los principios SRP, OCP y DIP se aplican de forma concreta en la separación de app.py, data.py, calc.py y persist.py.',
    'Los patrones Repository y Factory responden a necesidades puntuales (persistencia desacoplada y creación variada de usuarios), evitando su uso artificial.',
    'El prototipo funcional demuestra la viabilidad técnica de la propuesta y constituye la línea base para las siguientes etapas del Taller de Programación.',
    'La interfaz con Tailwind CSS permite un diseño responsivo sin necesidad de herramientas de compilación adicionales.',
    'El sistema permite la carga de evidencia fotográfica para incidentes disciplinarios, brindando trazabilidad verificable.',
]
for c in conclusiones:
    bullet(c)

# ── 18. Referencias ───────────────────────────────────────────────────
h1('18. Referencias bibliográficas')
refs = [
    'Evans, E. (2003). Domain-driven design: Tackling complexity in the heart of software. Addison-Wesley.',
    'Flask Documentation. (s.f.). Flask: Welcome to Flask. https://flask.palletsprojects.com/',
    'Gamma, E., Helm, R., Johnson, R. & Vlissides, J. (1994). Design patterns: Elements of reusable object-oriented software. Addison-Wesley.',
    'Martin, R. C. (2018). Clean architecture: A craftsman\'s guide to software structure and design. Prentice Hall.',
    'Python Documentation. (s.f.). Data structures — dataclasses. https://docs.python.org/3/library/dataclasses.html',
    'Tailwind CSS. (s.f.). Tailwind CSS: Rapidly build modern websites without ever leaving your HTML. https://tailwindcss.com/',
    'Werkzeug Documentation. (s.f.). Werkzeug: The Python WSGI Utility Library. https://werkzeug.palletsprojects.com/',
]
for ref in refs:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1.27)
    p.paragraph_format.first_line_indent = Cm(-1.27)
    p.paragraph_format.line_spacing = 2.0
    p.add_run(ref)

# ── Guardar ───────────────────────────────────────────────────────────
out = os.path.join(os.path.dirname(__file__), 'PROPUESTA_ScoreUp_APA.docx')
doc.save(out)
print(f'Documento generado: {out}')
